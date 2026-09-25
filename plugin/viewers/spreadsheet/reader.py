# Spreadsheet: a file-type viewer, written once for every corpus. A csv or tsv file is one sheet whose rows are its
# lines, cited as <path>#L<n> the way any text line is. An xlsx workbook is read with openpyxl (the kernel's environment
# needs it installed) and cited in the workbook's own notation, <path>#<Sheet>!<A1> for a cell, <Sheet>!<A1>:<B2> for a
# range and <Sheet> for a whole sheet, with a sheet name that holds a space or a # written with %20 or %23. The index
# keeps a delimited file's line offsets and a workbook's sheet names and sizes; cell values are read when asked for.
import csv
import io
import re
from urllib.parse import unquote

PAGE_ROWS = 200  # rows one sheet query returns
CELL = re.compile(r"^([A-Za-z]{1,3})(\d+)$")


def _kind(path):
    low = path.lower()
    return "xlsx" if low.endswith(".xlsx") else "tsv" if low.endswith(".tsv") else "csv"


def _open_book(path):
    import openpyxl  # a workbook needs openpyxl in the kernel's environment

    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def build_index(paths):
    """{path: {"kind", "offsets"? (a delimited file's byte offset of line n at n-1), "sheets": [{"name", "rows", "cols"?}]}}."""
    index = {}
    for path in paths:
        kind = _kind(path)
        if kind == "xlsx":
            try:
                book = _open_book(path)
            except Exception as e:  # noqa: BLE001 — a workbook that does not open is listed with the reason
                index[path] = {"kind": kind, "sheets": [], "error": f"{type(e).__name__}: {e}"}
                continue
            sheets = [{"name": ws.title, "rows": ws.max_row or 0, "cols": ws.max_column or 0} for ws in book.worksheets]
            book.close()
            index[path] = {"kind": kind, "sheets": sheets}
            continue
        offsets, pos = [], 0
        with open(path, "rb") as f:
            for raw in f:
                offsets.append(pos)
                pos += len(raw)
        index[path] = {"kind": kind, "offsets": offsets, "sheets": [{"name": path.rsplit("/", 1)[-1], "rows": len(offsets)}]}
    return index


def col_name(i):
    """0 -> A, 25 -> Z, 26 -> AA"""
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def col_index(name):
    n = 0
    for ch in name.upper():
        n = n * 26 + ord(ch) - 64
    return n - 1


def _delimited_rows(path, entry, start, count):
    """Rows start..start+count-1 (1-based line numbers) of a delimited file, each [line number, raw text, cells]."""
    out = []
    with open(path, "rb") as f:
        for n in range(start, min(len(entry["offsets"]), start + count - 1) + 1):
            f.seek(entry["offsets"][n - 1])
            raw = f.readline().decode("utf-8", "replace").rstrip("\r\n")
            cells = next(csv.reader(io.StringIO(raw), delimiter="\t" if entry["kind"] == "tsv" else ","), [])
            out.append([n, raw, cells])
    return out


def _sheet_rows(path, sheet, start, count):
    book = _open_book(path)
    try:
        ws = book[sheet]
        rows = []
        for n, row in enumerate(ws.iter_rows(min_row=start, max_row=start + count - 1, values_only=True), start):
            rows.append([n, ["" if v is None else v for v in row]])
        return rows
    finally:
        book.close()


def records(index, query):
    """{op: files} lists the files and their sheets; {op: sheet, path, sheet?, start?} gives PAGE_ROWS rows of one
    sheet from row `start`, each [row number, cells]."""
    query = query or {}
    if query.get("op") == "sheet":
        path = query.get("path")
        entry = index.get(path)
        if entry is None:
            return None
        start = max(1, int(query.get("start") or 1))
        if entry["kind"] == "xlsx":
            sheet = query.get("sheet") or (entry["sheets"][0]["name"] if entry["sheets"] else None)
            info = next((s for s in entry["sheets"] if s["name"] == sheet), None)
            if info is None:
                return {"path": path, "error": entry.get("error") or f"no sheet {sheet}"}
            rows = _sheet_rows(path, sheet, start, PAGE_ROWS)
            return {"path": path, "sheet": sheet, "sheets": [s["name"] for s in entry["sheets"]], "total": info["rows"],
                    "start": start, "rows": [[n, [str(v) for v in cells]] for n, cells in rows]}
        rows = _delimited_rows(path, entry, start, PAGE_ROWS)
        return {"path": path, "sheet": entry["sheets"][0]["name"], "sheets": [], "total": len(entry["offsets"]), "start": start,
                "rows": [[n, cells] for n, _, cells in rows]}
    return [{"path": p, "kind": e["kind"], "sheets": [s["name"] for s in e["sheets"]], "error": e.get("error")} for p, e in index.items()]


def resolve(index, locator):
    """A row of a delimited file (its excerpt the line as the file holds it) or a cell, a range or a sheet of a workbook
    (its excerpt the value, the values tab-separated per row, or the sheet's first row)."""
    path, fragment = locator.get("path"), str(locator.get("fragment") or "")
    entry = index.get(path)
    if entry is None or "key" in locator:
        return None
    if entry["kind"] != "xlsx":
        m = re.fullmatch(r"L(\d+)", fragment)
        if not m or not 1 <= int(m.group(1)) <= len(entry["offsets"]):
            return None
        n = int(m.group(1))
        (_, raw, _), = _delimited_rows(path, entry, n, 1)
        return {"excerpt": raw, "label": f"row {n}", "refs": [f"{path}#L{n}"], "target": {"path": path, "row": n}}
    sheet, _, cells = fragment.partition("!")
    sheet = unquote(sheet)
    info = next((s for s in entry["sheets"] if s["name"] == sheet), None)
    if info is None:
        return None
    ref = f"{path}#{fragment}"
    if not cells:
        first = _sheet_rows(path, sheet, 1, 1)
        return {"excerpt": "\t".join(str(v) for v in (first[0][1] if first else [])), "label": sheet[:30], "refs": [ref],
                "target": {"path": path, "sheet": sheet}}
    a, _, b = cells.partition(":")
    ma, mb = CELL.match(a), CELL.match(b or a)
    if not ma or not mb:
        return None
    r0, r1 = sorted((int(ma.group(2)), int(mb.group(2))))
    c0, c1 = sorted((col_index(ma.group(1)), col_index(mb.group(1))))
    rows = _sheet_rows(path, sheet, r0, r1 - r0 + 1)
    grid = [[str(row[c]) if c < len(row) else "" for c in range(c0, c1 + 1)] for _, row in rows]
    excerpt = "\n".join("\t".join(r) for r in grid)
    label = f"{sheet}!{cells} = {grid[0][0]}" if not b else f"{sheet}!{cells}"
    return {"excerpt": excerpt, "label": label[:40], "refs": [ref],
            "target": {"path": path, "sheet": sheet, "cell": [r0, c0], "to": [r1, c1]}}
